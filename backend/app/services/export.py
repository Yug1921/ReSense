"""
Turn a saved summary into a downloadable PDF or Word file.

Summaries are light Markdown written by the AI (headings, "**Label**" lines,
numbered labels, bullet and numbered lists, bold, italic, code). The same
reading rules as the app's own renderer (frontend rich-text.tsx) split it
into blocks, so the files look like the page, and anything unrecognised
stays a plain paragraph, so nothing is ever dropped. Layout flows, so it
works for a two-paragraph summary and a ten-page figure analysis alike.

Both formats use libraries already required by the backend: PyMuPDF
(its Story layout engine) for PDF and python-docx for Word.
"""
from __future__ import annotations

import html
import io
import re
from dataclasses import dataclass, field
from datetime import date

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

TONE_LABELS = {
    "academic": "Academic summary",
    "technical": "Technical summary",
    "research_insight": "Research Insight summary",
    "figure_table_analysis": "Figures and tables",
    "critical_review": "Critical review",
}

_HEADING_MD = re.compile(r"^#{1,4}\s+(.+?)\s*#*$")
_HEADING_BOLD_LINE = re.compile(r"^\*\*(.+?)\*\*:?$")
_HEADING_NUMBERED_LABEL = re.compile(r"^\d+[.)]\s+(?:\*\*)?([A-Z][^.!?]{1,80}?)(?:\*\*)?:?$")
_BULLET = re.compile(r"^[-*•]\s+(.*)$")
_ORDERED = re.compile(r"^\d+[.)]\s+(.*)$")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


@dataclass
class Block:
    kind: str  # heading | paragraph | list
    text: str = ""
    ordered: bool = False
    items: list[str] = field(default_factory=list)


def parse_blocks(source: str) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []
    current: Block | None = None

    def flush_paragraph():
        nonlocal paragraph
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph)))
        paragraph = []

    def flush_list():
        nonlocal current
        if current:
            blocks.append(current)
        current = None

    for raw in source.replace("\r", "").split("\n"):
        line = raw.strip()
        if not line:
            flush_paragraph()
            flush_list()
            continue
        match = _HEADING_MD.match(line) or _HEADING_BOLD_LINE.match(line) or _HEADING_NUMBERED_LABEL.match(line)
        if match:
            flush_paragraph()
            flush_list()
            text = re.sub(r"^\d+[.)]\s+", "", match.group(1).replace("**", "")).rstrip(":").strip()
            blocks.append(Block("heading", text))
            continue
        item = _BULLET.match(line)
        ordered = None if item else _ORDERED.match(line)
        if item or ordered:
            flush_paragraph()
            if not current or current.ordered != bool(ordered):
                flush_list()
                current = Block("list", ordered=bool(ordered))
            current.items.append((item or ordered).group(1))
            continue
        if current:
            current.items[-1] += f" {line}"  # wrapped continuation of a list item
            continue
        paragraph.append(line)
    flush_paragraph()
    flush_list()
    return blocks


def _runs(text: str) -> list[tuple[str, str]]:
    """Split inline Markdown into (style, text) runs: plain, bold, italic, code."""
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append(("bold", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append(("code", part[1:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append(("italic", part[1:-1]))
        else:
            runs.append(("plain", part))
    return runs


def _title(filename: str) -> str:
    name = re.sub(r"\.(pdf|docx?)$", "", filename or "Research paper", flags=re.IGNORECASE)
    return re.sub(r"[_]+", " ", name).strip() or "Research paper"


def download_name(filename: str, tone: str, extension: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "-", _title(filename)).strip("-")[:60] or "paper"
    return f"{base}-{tone.replace('_', '-')}.{extension}"


# ---------------------------------------------------------------- PDF

_PDF_CSS = """
* { font-family: sans-serif; }
body { font-size: 10.5pt; line-height: 1.5; color: #222226; }
h1 { font-size: 19pt; line-height: 1.25; color: #111114; margin: 0 0 4pt 0; }
.meta { font-size: 9pt; color: #6a6a71; margin: 0 0 16pt 0; }
h2 { font-size: 12.5pt; color: #2f3794; margin: 14pt 0 5pt 0; }
p { margin: 0 0 8pt 0; text-align: justify; }
ul, ol { margin: 0 0 8pt 0; }
li { margin: 0 0 3pt 0; }
code { font-family: monospace; font-size: 9.5pt; }
"""


def _inline_html(text: str) -> str:
    out = []
    for style, part in _runs(text):
        safe = html.escape(part)
        out.append({"bold": f"<b>{safe}</b>", "italic": f"<i>{safe}</i>", "code": f"<code>{safe}</code>"}.get(style, safe))
    return "".join(out)


def _summary_html(title: str, subtitle: str, blocks: list[Block]) -> str:
    parts = [f"<h1>{html.escape(title)}</h1>", f'<p class="meta">{html.escape(subtitle)}</p>']
    for block in blocks:
        if block.kind == "heading":
            parts.append(f"<h2>{_inline_html(block.text)}</h2>")
        elif block.kind == "list":
            tag = "ol" if block.ordered else "ul"
            items = "".join(f"<li>{_inline_html(item)}</li>" for item in block.items)
            parts.append(f"<{tag}>{items}</{tag}>")
        else:
            parts.append(f"<p>{_inline_html(block.text)}</p>")
    return "<body>" + "".join(parts) + "</body>"


def summary_pdf(filename: str, tone: str, content: str) -> bytes:
    title = _title(filename)
    subtitle = f"{TONE_LABELS.get(tone, 'Summary')} · Written by ReSense from the paper's text · {date.today():%d %B %Y}"
    story = fitz.Story(html=_summary_html(title, subtitle, parse_blocks(content)), user_css=_PDF_CSS)
    page = fitz.paper_rect("a4")
    body = page + (56, 56, -56, -64)  # margins: left, top, right, bottom

    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(body)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Footer on every page: title on the left, page numbers on the right.
    document = fitz.open("pdf", buffer.getvalue())
    total = document.page_count
    for number, pdf_page in enumerate(document, 1):
        y = page.height - 34
        pdf_page.insert_text((56, y), title[:80], fontsize=8, color=(0.45, 0.45, 0.48))
        label = f"Page {number} of {total}"
        width = fitz.get_text_length(label, fontsize=8)
        pdf_page.insert_text((page.width - 56 - width, y), label, fontsize=8, color=(0.45, 0.45, 0.48))
    document.set_metadata({"title": f"{title}: {TONE_LABELS.get(tone, 'Summary')}", "creator": "ReSense"})
    return document.tobytes(garbage=3, deflate=True)


# ---------------------------------------------------------------- Word

_ACCENT = RGBColor(0x2F, 0x37, 0x94)
_MUTED = RGBColor(0x6A, 0x6A, 0x71)


def _add_runs(paragraph, text: str) -> None:
    for style, part in _runs(text):
        run = paragraph.add_run(part)
        run.bold = style == "bold"
        run.italic = style == "italic"
        if style == "code":
            run.font.name = "Consolas"


def summary_docx(filename: str, tone: str, content: str) -> bytes:
    title = _title(filename)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = f"{title} · {TONE_LABELS.get(tone, 'Summary')} · ReSense"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = _MUTED

    heading = document.add_heading(title, level=0)
    heading.runs[0].font.size = Pt(22)
    meta = document.add_paragraph()
    meta_run = meta.add_run(
        f"{TONE_LABELS.get(tone, 'Summary')} · Written by ReSense from the paper's text · {date.today():%d %B %Y}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = _MUTED

    for block in parse_blocks(content):
        if block.kind == "heading":
            h = document.add_heading(level=2)
            _add_runs(h, block.text)
            for run in h.runs:
                run.font.color.rgb = _ACCENT
        elif block.kind == "list":
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                _add_runs(document.add_paragraph(style=style), item)
        else:
            _add_runs(document.add_paragraph(), block.text)

    document.core_properties.title = f"{title}: {TONE_LABELS.get(tone, 'Summary')}"
    document.core_properties.author = "ReSense"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
